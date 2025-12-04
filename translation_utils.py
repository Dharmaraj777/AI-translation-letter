import json
import os
from io import BytesIO
from typing import Iterable, List, Dict, Any

import fitz  # PyMuPDF
from docx import Document
from pptx import Presentation
from azure.storage.blob import BlobClient

from spanish_translator_logger import logger
from spanish_translator_config_loader import ConfigLoader



# UtilityFunctions 
# -----------------------------
class UtilityFunctions:

    def __init__(self):
        cfg = ConfigLoader.get_instance()
        # Use the container client that ConfigLoader already exposes
        self.input_container_client = cfg.input_container_client
        # If you later want to upload logs from here:
        self.logs_container_client = getattr(cfg, "logs_container_client", None)

    def get_files_to_process(self) -> List[str]:
        """
        Return all blobs in the input container.
        If you want 'today only', you can filter with blob.last_modified.date().
        """
        blob_names: List[str] = []
        for blob in self.input_container_client.list_blobs():
            blob_names.append(blob.name)
        return blob_names

    def get_page_count_from_blob(self, blob_name: str) -> int:
        blob_client: BlobClient = self.input_container_client.get_blob_client(blob_name)
        file_data = blob_client.download_blob().readall()

        _, ext = os.path.splitext(blob_name)
        ext = ext.lower()

        if ext == ".pdf":
            with fitz.open(stream=file_data, filetype="pdf") as doc:
                return len(doc)

        elif ext == ".pptx":
            prs = Presentation(BytesIO(file_data))
            return len(prs.slides)

        elif ext == ".docx":
            document = Document(BytesIO(file_data))
            para_count = len(document.paragraphs)
            # crude heuristic: ~40 paragraphs per "page"
            approx_pages = max(1, para_count // 40 or 1)
            return approx_pages

        else:
            raise ValueError(f"Unsupported file type for page count: {ext}")

    def download_blob_bytes(self, blob_name: str) -> bytes:
        blob_client: BlobClient = self.input_container_client.get_blob_client(blob_name)
        logger.info(f"Downloading blob: {blob_name}")
        return blob_client.download_blob().readall()

    def upload_log_to_blob(self, blob_name: str, log_data: str) -> None:
        if not self.logs_container_client:
            logger.warning("logs_container_client not configured; cannot upload log blob.")
            return

        try:
            blob_client = self.logs_container_client.get_blob_client(blob_name)
            blob_client.upload_blob(log_data.encode("utf-8"), overwrite=True)
            logger.info(f"Log file {blob_name} uploaded successfully.")
        except Exception as e:
            logger.error(f"Failed to upload log file {blob_name}: {str(e)}")
